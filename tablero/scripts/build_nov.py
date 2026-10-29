import sys, datetime as dt, openpyxl
from openpyxl.comments import Comment
src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src, keep_vba=True)
ws = wb['GRANITE ']
# drop merges below the data block, re-add shifted after deleting rows
low = [m for m in list(ws.merged_cells.ranges) if m.min_row >= 19]
for m in low: ws.unmerge_cells(str(m))
ws.delete_rows(15, 2)   # UGC_Trybe, old RMKT row
ws.delete_rows(8, 2)    # Amazon traffic + its subtotal
for m in low:
    ws.merge_cells(start_row=m.min_row-4, start_column=m.min_col, end_row=m.max_row-4, end_column=m.max_col)
ws.row_dimensions[13].height = 14.25
ws.row_dimensions[17].height = None

S, E = dt.datetime(2026,11,1), dt.datetime(2026,11,30)
lines = [
 # F, G, H, I, J, Q, R, T
 ('Campaña de marca- Exact match (Brand Granite)', 'Google', 'SEM', 'Sales', 'Text', 0.08, 430, 22),
 ('Campaña de Performance Shopping (Granite_Shopping)', 'Google', 'Shopping', 'Sales', 'Merchant', 0.02, 800, 94),
 ('G5_RMKT_DemandGen_Visitors30d: remarketing visitantes 30 días sin compra (Proteína, Recovery, EAA)', 'Google', 'Demand Gen RMKT', 'Sales', 'Imagen - video ', 0.01, 450, 15),
 ('Advantage+ Landing conversión Creatina, The meadows Whey, recovery - ROAS OBJETIVO', 'Meta  ', 'Facebook - Instagram', 'Sales', 'imagen + video + catálogo', 0.01, 940, 18.03),
 ('Granite_RMKT_Dinamico_VC-ATC-Visit_NoCompra', 'Meta  ', 'Facebook - Instagram', 'Sales', 'imagen + video + catálogo', 0.01, 568, 18.03),
]
for i, (f, g, h, obj, fmt, q, r, t) in enumerate(lines):
    n = 8 + i
    ws[f'B{n}'], ws[f'C{n}'] = S, E
    ws[f'D{n}'] = f'=C{n}-B{n}+1'
    ws[f'E{n}'] = 'Performance'
    for col, v in zip('FGHIJ', (f, g, h, obj, fmt)): ws[f'{col}{n}'] = v
    ws[f'K{n}'], ws[f'L{n}'] = 'CPA', 'Purchase'
    ws[f'M{n}'] = f'=O{n}/Q{n}'; ws[f'N{n}'] = f'=M{n}/$N$5'
    ws[f'O{n}'] = f'=P{n}/3%';   ws[f'P{n}'] = f'=R{n}/T{n}'
    ws[f'Q{n}'], ws[f'R{n}'], ws[f'T{n}'] = q, r, t
    ws[f'S{n}'] = f'=R{n}/$F$5'
for c in 'MNOPRS': ws[f'{c}13'] = f'=SUM({c}8:{c}12)'
for c in 'MNOPQRS': ws[f'{c}14'] = f'={c}13'
ws['T14'] = '=R14/P14'
ws['D16'] = '=+R14'; ws['R16'] = '=R15+R14'
ws['S18'] = '=R18*R15'; ws['S19'] = '=R14*R18'; ws['S20'] = '=S18+S19'
ws['I4'] = 0; ws['J4'] = '=R13'
ws['F5'] = '=R14'
ws['R11'].comment = Comment('Nov: 640 de oct + 300 reasignados (Meta Amazon 204 + UGC_Trybe 96).', 'Rebold')
ws['R12'].comment = Comment('Nov: 313 de oct + 255 reasignados (UGC_Trybe 67 + Tráfico Amazon 188).', 'Rebold')
ws['R10'].comment = Comment('Reemplaza Demand Gen Categorización audiencias (450, pausada 26-oct).', 'Rebold')
rs = wb['Resume']
rs['E12'] = "='GRANITE '!R15"; rs['F12'] = "='GRANITE '!M14"; rs['G12'] = "='GRANITE '!O14"
rs['D16'] = "='GRANITE '!R14"; rs['D17'] = "='GRANITE '!R15"
wb.save(out)
